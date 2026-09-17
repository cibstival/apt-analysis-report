#!/usr/bin/env python3
"""생성된 엑셀의 수식 오류 검사.

LibreOffice(soffice)가 있으면 수식을 계산한 사본을 만들어 #DIV/0!, #N/A, #NAME?, #REF!, #VALUE! 셀을 찾는다.
없으면 건너뛴다(엑셀에서 열면 자동 계산됨).
사용법: python scripts/verify_xlsx.py output/<파일>.xlsx
"""
import shutil, subprocess, sys, tempfile
from pathlib import Path
from openpyxl import load_workbook

p = Path(sys.argv[1]).resolve()
soffice = shutil.which("soffice") or shutil.which("libreoffice")
if not soffice:
    print("LibreOffice 없음 → 수식 검사 생략 (엑셀에서 열어 확인)"); sys.exit(0)
with tempfile.TemporaryDirectory() as td:
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", td, str(p)],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(Path(td) / p.name, data_only=True)
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") and c.value.rstrip("!?").upper() in ("#DIV/0", "#N/A", "#NAME", "#REF", "#VALUE", "#NUM", "#NULL"):
                    errs.append(f"{ws.title}!{c.coordinate}={c.value}")
    if errs:
        print(f"수식 오류 {len(errs)}개:"); print("\n".join(errs[:50])); sys.exit(1)
    print("수식 오류 없음")
