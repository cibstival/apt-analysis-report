#!/usr/bin/env python3
"""아파트 단지 분석 엑셀 보고서 생성기.

사용법:
  python scripts/build_report.py --apt config/apartments/<단지>.json [--market config/market.json] [--out output/<파일>.xlsx]

시트 구성(고정):
  요약_차트 · <단지>_분석 · 데이터 · 거래량 · <단지>_실거래 · 금리변경이력 · 출처_가정
모든 계산 셀은 엑셀 수식이며, 분석 문장은 아파트 설정 JSON의 text 항목에서 가져온다.
"""
import argparse, math, sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib_report import (PY_M2, acquisition_tax_formula, brokerage_formula, compute_metrics, fill, find_unfilled, half_key, load_json, load_policy, load_trades,
                        parse_date, safe_sheet)

F = "Arial"
def font(**k): return Font(name=F, **k)
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill("solid", fgColor="1F3864"); YEL = PatternFill("solid", fgColor="FFF2CC")
BLUE = Font(name=F, color="0000FF"); GREEN = Font(name=F, color="008000")
NAVY = "1F3864"; BLUE2 = "2F5496"; LBLUE = "D9E1F2"; PALE = "EAF1FB"; CREAM = "FFF2CC"; GREY = "F2F2F2"; ORANGE = "C65911"
GREENF = "E2EFDA"; REDF = "FCE4D6"
def FILL(c): return PatternFill("solid", fgColor=c)


def header(ws, row, cols_titles, height=34):
    for col, t in cols_titles:
        c = ws.cell(row=row, column=col, value=t); c.font = font(bold=True, color="FFFFFF"); c.fill = HDR; c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


def style_line(s, color, width=38100, dash=None, marker="circle"):
    s.graphicalProperties.line.solidFill = color; s.graphicalProperties.line.width = width
    if dash: s.graphicalProperties.line.dashStyle = dash
    s.marker.symbol = marker; s.marker.size = 7
    s.marker.graphicalProperties.solidFill = color; s.marker.graphicalProperties.line.solidFill = color; s.smooth = False


def secondary(main, sec):
    for ax in (main.x_axis, main.y_axis, sec.y_axis): ax.delete = False
    sec.y_axis.axId = 200; sec.y_axis.crosses = "max"; sec.y_axis.majorGridlines = None
    sec.x_axis.axId = 500; sec.x_axis.delete = True; sec.y_axis.crossAx = 500; sec.x_axis.crossAx = 200
    main += sec


def nice_max(v, step):
    return max(step, math.ceil(v * 1.15 / step) * step)


def policy_rows():
    """출처_가정 시트용 정책 기준 행. config/policy.json의 세제·규제·대출 항목과 확인일을 그대로 옮긴다."""
    pol = load_policy(); out = []
    for key in ("acquisition_tax", "brokerage", "capital_gains", "loan", "regulation", "tax_reform"):
        sec = pol.get(key)
        if not sec:
            continue
        parts = []
        if key == "acquisition_tax":
            parts.append(f"{sec['low_upto']/10000:g}억 이하 {sec['low_rate']:.0%}, {sec['high_from']/10000:g}억 초과 {sec['high_rate']:.0%}, {sec['interp_note']}. {sec['extra_note']}")
        elif key == "brokerage":
            parts.append(" / ".join((f"~{b['upto']/10000:g}억" if b["upto"] else "초과") + f" {b['rate']:.1%}" for b in sec["brackets"]) + f" + 부가세 {sec['vat_ratio']:.0%}. {sec['sell_note']}")
        elif key == "capital_gains":
            parts.append(sec["assumption"])
        for k, v in (sec.get("current") or {}).items():
            parts.append(f"{k}: {v}")
        if sec.get("status"):
            parts.append(f"[{sec['status']}]")
        parts.append(f"— 출처 {sec.get('source', '')}, 시행 {sec.get('effective_from', '')}~")
        out.append((f"정책 기준: {sec.get('label', key)}", " ".join(parts)))
    out.append(("정책 기준 확인일", f"{pol['checked_at']} (config/policy.json, {pol['check_interval_days']}일 경과 시 실행 전 재확인). 최근 변경: " + "; ".join(f"{c['date']} {c['note']}" for c in pol.get("changelog", [])[-3:])))
    return out


def build(apt_path, mkt_path, out_path):
    base = Path(apt_path).resolve().parent.parent.parent  # 스킬 루트
    apt = load_json(apt_path); mkt = load_json(mkt_path); apt["_config_dir"] = str(Path(apt_path).resolve().parent)
    trades = load_trades(apt, base)
    M = compute_metrics(apt, mkt, trades)
    TXT = fill(apt.get("text", {}), M)
    left = find_unfilled(TXT)
    if left:
        print("[경고] 치환되지 않은 플레이스홀더:", left[:10])

    NAME = apt["name"]; SHORT = apt.get("sheet_prefix", NAME)
    SN_T = safe_sheet(SHORT, "_실거래"); SN_A = safe_sheet(SHORT, "_분석")
    types = apt["types"]; focus = apt["user_deal"]["pyeong"]; comp = apt.get("compare_pyeong")
    G = mkt["groups"]; HN, LN = G["high"]["name"], G["low"]["name"]
    asof = parse_date(mkt["as_of"])
    periods = [(p["label"], parse_date(p["date"]), p.get("partial", False)) for p in mkt["periods"]]
    first_lab, last_lab = periods[0][0], periods[-1][0]

    wb = Workbook()
    ws = wb.active; ws.title = "요약_차트"
    wd = wb.create_sheet("데이터"); wv = wb.create_sheet("거래량"); wt = wb.create_sheet(SN_T)
    we = wb.create_sheet("금리변경이력"); wn = wb.create_sheet("출처_가정")

    # ===================== 금리변경이력 =====================
    kr, us = mkt["rates"]["kr"], mkt["rates"]["us"]
    we["A1"] = f"기준금리 변경 이력 ({kr[0]['date'][:7]} ~ {mkt['as_of'][:7]})"; we["A1"].font = font(bold=True, size=14)
    we["A2"] = "파란 글씨 = 공표된 실제 결정값. '데이터' 시트의 반기말 금리는 이 표에서 INDEX/MATCH로 자동 조회됩니다."
    we["A2"].font = font(italic=True, color="595959", size=9)
    header(we, 4, [(1, "한국은행 결정일"), (2, "기준금리"), (3, "변경폭/비고"), (5, "FOMC 결정일"), (6, "FFR 목표 상단"), (7, "목표범위/비고")], 20)
    for col0, evs in ((1, kr), (5, us)):
        for i, e in enumerate(evs):
            r = 5 + i
            we.cell(r, col0, parse_date(e["date"])).number_format = "yyyy-mm-dd"
            we.cell(r, col0 + 1, e["rate"]).number_format = "0.00%"; we.cell(r, col0 + 2, e.get("note", ""))
            for c in (col0, col0 + 1): we.cell(r, c).font = BLUE
            we.cell(r, col0 + 2).font = font()
            for c in range(col0, col0 + 3): we.cell(r, c).border = box
    KR_END = 4 + len(kr); US_END = 4 + len(us)
    for c, w in zip("ABCDEFG", [16, 12, 30, 3, 16, 14, 32]): we.column_dimensions[c].width = w

    # ===================== 실거래 =====================
    wt["A1"] = f"{NAME} ({apt.get('address','')}) 매매 실거래 내역 {trades[0]['date']:%Y.%m}~{trades[-1]['date']:%Y.%m}"
    wt["A1"].font = font(bold=True, size=14)
    wt["A2"] = apt.get("trades_note", "출처: 국토교통부 실거래가. C열 '동'이 공란인 거래는 동 정보 미확인. I열 0 = 분석 제외(이상거래).")
    wt["A2"].font = font(italic=True, color="C00000", size=9)
    header(wt, 4, [(1, "계약일"), (2, "반기키"), (3, "동\n(입력)"), (4, "공급평형\n(평)"), (5, "전용면적\n(㎡)"), (6, "층"),
                   (7, "거래가\n(만원)"), (8, "전용 평단가\n(만원/3.3㎡)"), (9, "포함\n(1/0)"), (10, "비고")])
    TS = 5; NT = len(types); TY0 = 6; TY1 = TY0 + NT - 1; WROW = TY1 + 1
    for i, t in enumerate(trades):
        r = TS + i
        wt.cell(r, 1, t["date"]).number_format = "yyyy-mm-dd"
        wt.cell(r, 2, f'=YEAR(A{r})&"."&IF(MONTH(A{r})<=6,"06","12")')
        wt.cell(r, 3, t["dong"] or None); wt.cell(r, 3).fill = YEL
        wt.cell(r, 4, t["pyeong"]); wt.cell(r, 6, t["floor"]); wt.cell(r, 7, t["price"]).number_format = "#,##0"
        wt.cell(r, 5, f"=INDEX($N${TY0}:$N${TY1},MATCH(D{r},$M${TY0}:$M${TY1},0))").number_format = "0.00"
        wt.cell(r, 8, f"=G{r}/(E{r}/3.3058)").number_format = "#,##0"
        wt.cell(r, 9, t["include"]); wt.cell(r, 9).fill = YEL
        wt.cell(r, 10, t["note"] or None)
        for c in (1, 4, 6, 7, 9): wt.cell(r, c).font = BLUE
        for c in range(1, 11):
            wt.cell(r, c).border = box
            if wt.cell(r, c).font.name != F: wt.cell(r, c).font = font()
    TE = TS + len(trades) - 1
    wt["M4"] = "평형 정보"; wt["M4"].font = font(bold=True)
    kb_asof = apt.get("kb_asof", "")
    header(wt, 5, [(13, "공급평형"), (14, "전용㎡"), (15, "세대수"), (16, f"KB 매매시세\n(만원, {kb_asof})"), (17, "전용 평단가")], 30)
    TYROW = {}
    for i, ty in enumerate(types):
        r = TY0 + i; TYROW[ty["pyeong"]] = r
        for c, v in zip((13, 14, 15, 16), (ty["pyeong"], ty["area"], ty.get("households", 0), ty.get("kb_price"))):
            wt.cell(r, c, v).font = BLUE; wt.cell(r, c).border = box
        wt.cell(r, 16).number_format = "#,##0"
        wt.cell(r, 17, f'=IF(P{r}="","",P{r}/(N{r}/3.3058))').number_format = "#,##0"; wt.cell(r, 17).border = box; wt.cell(r, 17).font = font()
    wt.cell(WROW, 13, "세대가중 평단가").font = font(bold=True)
    wt.cell(WROW, 17, f'=IFERROR(SUMPRODUCT(O{TY0}:O{TY1},Q{TY0}:Q{TY1})/SUM(O{TY0}:O{TY1}),"")').number_format = "#,##0"  # 세대수·KB시세 미입력 시 공란
    wt.cell(WROW, 17).font = font(bold=True)
    wt.cell(WROW + 1, 13, "※ 마지막(부분) 반기에 신고된 실거래가 없으면 차트 마지막 점은 이 KB시세 세대가중 평단가를 사용").font = font(size=9, color="595959")
    for c, w in zip("ABCDEFGHIJKLMNOPQ", [12, 9, 7, 9, 9, 6, 11, 12, 7, 44, 3, 3, 10, 9, 8, 13, 11]): wt.column_dimensions[c].width = w
    wt.freeze_panes = "A5"
    R = lambda col: f"'{SN_T}'!${col}${TS}:${col}${TE}"
    g_, d_, inc_, key_, dt_, dong_, ppp_ = R("G"), R("D"), R("I"), R("B"), R("A"), R("C"), R("H")

    # ===================== 데이터 =====================
    wd["A1"] = "반기별 데이터: 평균 평단가(만원/3.3㎡, 전용면적 기준) vs 한·미 기준금리"; wd["A1"].font = font(bold=True, size=14)
    wd["A2"] = f"노란 칸 = 입력값. C·D열(구 그룹 평단가)은 공표치 기준 추정치. E~G열({NAME})은 실거래 시트에서 자동 계산."
    wd["A2"].font = font(italic=True, color="C00000", size=9)
    wd["I3"] = f"{NAME} 분석대상 동:"; wd["I3"].font = font(bold=True); wd["I3"].alignment = Alignment(horizontal="right")
    wd["K3"] = "전체"; wd["K3"].fill = YEL; wd["K3"].font = font(bold=True, color="0000FF"); wd["K3"].alignment = Alignment(horizontal="center")
    wd["K3"].comment = Comment("'전체' 또는 동 번호(예: 101). 실거래 시트 C열에 해당 동이 입력된 거래만 평균합니다.", "report")
    header(wd, 4, [(1, "반기말 시점"), (2, "기준일(조회용)"), (3, f"{HN}\n평균 평단가"), (4, f"{LN}\n평균 평단가"), (5, f"{NAME}\n거래건수"),
                   (6, f"{NAME}\n실거래 평균"), (7, f"{NAME}\n평단가(차트)"), (8, "고가/저가\n배율"), (9, "한국\n기준금리"),
                   (10, "미국\n기준금리(상단)"), (11, "한·미\n금리차(한-미)"), (12, "비고")])
    S = 5; L = S + len(periods) - 1
    for i, (lab, d, partial) in enumerate(periods):
        r = S + i
        wd.cell(r, 1, lab); wd.cell(r, 2, d).number_format = "yyyy-mm-dd"
        for c, v in ((3, G["high"]["ppp"][i]), (4, G["low"]["ppp"][i])):
            wd.cell(r, c, v).font = BLUE; wd.cell(r, c).fill = YEL; wd.cell(r, c).number_format = "#,##0"
        hk = half_key(d)
        if partial:
            wd.cell(r, 5, f'=IF($K$3="전체",COUNTIFS({key_},"{hk}",{inc_},1),COUNTIFS({key_},"{hk}",{inc_},1,{dong_},$K$3))')
            wd.cell(r, 6, f'=IF(E{r}>0,IF($K$3="전체",AVERAGEIFS({ppp_},{key_},"{hk}",{inc_},1),AVERAGEIFS({ppp_},{key_},"{hk}",{inc_},1,{dong_},$K$3)),\'{SN_T}\'!$Q${WROW})')
            wd.cell(r, 7, f"=F{r}")
            wd.cell(r, 12, "부분 반기: 신고 실거래 없으면 KB시세 세대가중 평단가")
        else:
            wd.cell(r, 5, f'=IF($K$3="전체",COUNTIFS({key_},A{r},{inc_},1),COUNTIFS({key_},A{r},{inc_},1,{dong_},$K$3))')
            wd.cell(r, 6, f'=IF(E{r}=0,"",IF($K$3="전체",AVERAGEIFS({ppp_},{key_},A{r},{inc_},1),AVERAGEIFS({ppp_},{key_},A{r},{inc_},1,{dong_},$K$3)))')
            if i == 0:
                wd.cell(r, 7, f'=IF(F{r}="",IF(F{r+1}="",0,F{r+1}),F{r})')
            else:
                wd.cell(r, 7, f'=IF(F{r}<>"",F{r},IF(N(F{r+1})>0,(G{r-1}+F{r+1})/2,G{r-1}))')
        wd.cell(r, 8, f"=C{r}/D{r}").number_format = '0.00"배"'
        wd.cell(r, 9, f"=INDEX(금리변경이력!$B$5:$B${KR_END},MATCH(B{r},금리변경이력!$A$5:$A${KR_END},1))")
        wd.cell(r, 10, f"=INDEX(금리변경이력!$F$5:$F${US_END},MATCH(B{r},금리변경이력!$E$5:$E${US_END},1))")
        wd.cell(r, 11, f"=(I{r}-J{r})*100").number_format = '+0.00"%p";-0.00"%p";0.00"%p"'
        for c in (6, 7): wd.cell(r, c).number_format = "#,##0"
        for c in (9, 10): wd.cell(r, c).number_format = "0.00%"; wd.cell(r, c).font = GREEN
        wd.cell(r, 6).font = GREEN
        for grp, col in ((G["high"], 3), (G["low"], 4)):
            if lab in grp.get("anchors", {}):
                wd.cell(r, col).comment = Comment("근거 공표치(만원/3.3㎡): " + grp["anchors"][lab], "report")
        for c in range(1, 13):
            cell = wd.cell(r, c); cell.border = box
            if c < 12: cell.alignment = Alignment(horizontal="center")
            if cell.font.name != F: cell.font = font()
        wd.cell(r, 12).font = font(size=9, color="595959")
    small = [k for k, v in M["_halves"].items() if 0 < v["n"] <= 2 and k >= first_lab]
    zero = [lab for lab, d, p in periods if not p and M["_halves"].get(lab, {"n": 0})["n"] == 0]
    notes = [f"※ C·D열: 공표 기준점(셀 메모) 사이를 변동률 흐름에 맞춰 보간한 추정치. 그룹값 = 구별 단순평균({G['high']['members']} / {G['low']['members']}).",
             f"※ E~G열: {NAME} 반기 내 실거래의 전용 평단가 단순평균. 거래 1~2건 반기({', '.join(small) or '없음'})는 평형 구성에 따라 크게 튈 수 있음. 거래 없는 반기({', '.join(zero) or '없음'})는 앞뒤 평균으로 보간.",
             "※ 소형 평형은 전용 평단가가 대형보다 높게 나오므로, 그 반기에 소형 거래 비중이 높으면 평균이 올라갑니다."]
    for j, n in enumerate(notes):
        wd.cell(L + 2 + j, 1, n).font = font(size=9, color="595959")
    for c, w in zip("ABCDEFGHIJKL", [11, 12, 13, 12, 12, 13, 13, 10, 10, 12, 12, 40]): wd.column_dimensions[c].width = w
    wd.freeze_panes = "B5"

    # ===================== 거래량 =====================
    vol = mkt["volume"]["rows"]
    wv["A1"] = "반기별 아파트 매매 거래량과 전반기 대비 증감률"; wv["A1"].font = font(bold=True, size=14)
    wv["A2"] = f"노란 칸(C·D열) = 추정 입력값(서울 전체 추정 거래량 × 그룹 비중). E열({NAME})은 실거래 시트에서 자동 집계(이상거래 포함 전체 건수)."
    wv["A2"].font = font(italic=True, color="C00000", size=9)
    header(wv, 4, [(1, "반기"), (2, "서울 전체\n(참고·추정)"), (3, f"{HN}\n거래량"), (4, f"{LN}\n거래량"), (5, f"{NAME}\n거래량"),
                   (6, f"{HN}\n증감률"), (7, f"{LN}\n증감률"), (8, f"{NAME}\n증감률"), (9, "한국\n기준금리"), (10, "미국\n기준금리(상단)"), (11, "비고")])
    VS = 5; VL = VS + len(vol) - 1
    for i, v in enumerate(vol):
        r = VS + i; wv.cell(r, 1, v["label"])
        if v.get("seoul"):
            wv.cell(r, 2, v["seoul"]).font = BLUE; wv.cell(r, 2).number_format = "#,##0"
            for c, val in ((3, round(v["seoul"] * v["high_share"])), (4, round(v["seoul"] * v["low_share"]))):
                wv.cell(r, c, val).font = BLUE; wv.cell(r, c).fill = YEL; wv.cell(r, c).number_format = "#,##0"
            wv.cell(r, 5, f'=IF(데이터!$K$3="전체",COUNTIFS({key_},A{r}),COUNTIFS({key_},A{r},{dong_},데이터!$K$3))')
            if i > 0:
                for c, src in ((6, "C"), (7, "D"), (8, "E")):
                    wv.cell(r, c, f"=({src}{r}-{src}{r-1})/MAX({src}{r-1},1)").number_format = '+0%;-0%;0%'
            else:
                wv.cell(r, 11, "증감률 계산용 기준 반기")
        else:
            wv.cell(r, 11, "부분 반기: 신고 집계 중이라 반기 비교 불가 → 공란(차트에서 끊김)")
        if i > 0:
            dr = S + i - 1
            wv.cell(r, 9, f"=데이터!I{dr}").number_format = "0.00%"; wv.cell(r, 10, f"=데이터!J{dr}").number_format = "0.00%"
            wv.cell(r, 9).font = GREEN; wv.cell(r, 10).font = GREEN
        for c in range(1, 12):
            cell = wv.cell(r, c); cell.border = box
            if c < 11: cell.alignment = Alignment(horizontal="center")
            if cell.font.name != F: cell.font = font()
        wv.cell(r, 11).font = font(size=9, color="595959")
    wv.cell(VL + 2, 1, "※ 증감률 = (이번 반기 − 직전 반기) ÷ 직전 반기. 직전 반기가 0건이면 1건으로 나눔.").font = font(size=9, color="595959")
    wv.cell(VL + 3, 1, f"※ {NAME}는 반기 거래 표본이 작아 1~2건 차이로 증감률이 크게 움직입니다. 방향성 위주로 보세요.").font = font(size=9, color="595959")
    for c, w in zip("ABCDEFGHIJK", [10, 12, 13, 11, 13, 13, 11, 13, 10, 12, 58]): wv.column_dimensions[c].width = w
    wv.freeze_panes = "B5"

    # ===================== 요약_차트 =====================
    ws["A1"] = f"기준금리와 서울 아파트 평균 평단가: {HN} vs {LN} vs {NAME}, {first_lab}~{last_lab}"
    ws["A1"].font = font(bold=True, size=16, color=NAVY)
    ws["A2"] = "왼쪽 축: 평균 평단가(만원/3.3㎡, 전용면적 기준)  |  오른쪽 보조축: 한국·미국 기준금리  |  X축: 6개월(반기말) 단위, *는 부분 반기"
    ws["A2"].font = font(size=10, color="595959")
    ymax = nice_max(max(G["high"]["ppp"] + G["low"]["ppp"] + [h["avg_ppp"] or 0 for h in M["_halves"].values()]), 2000)
    rmax = max(e["rate"] for e in kr + us)
    ch = LineChart(); ch.title = "① 평균 평단가 vs 한·미 기준금리"; ch.height = 15; ch.width = 36
    ch.y_axis.title = "평균 평단가 (만원/3.3㎡)"; ch.y_axis.number_format = "#,##0"; ch.y_axis.majorGridlines = None
    ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = ymax; ch.y_axis.majorUnit = ymax / 6
    ch.x_axis.title = "반기말 시점 (6개월 단위)"
    for col in (3, 4, 7): ch.add_data(Reference(wd, min_col=col, min_row=S, max_row=L), titles_from_data=False)
    ch.set_categories(Reference(wd, min_col=1, min_row=S, max_row=L))
    for s, col, nm in zip(ch.series, ["C00000", "ED7D31", "7030A0"], [f"{HN} 평균 평단가", f"{LN} 평균 평단가", f"{NAME} 평균 평단가(실거래)"]):
        s.tx = SeriesLabel(v=nm); style_line(s, col)
    ch2 = LineChart(); ch2.add_data(Reference(wd, min_col=9, max_col=10, min_row=S, max_row=L), titles_from_data=False)
    ch2.y_axis.title = "기준금리 (보조축)"; ch2.y_axis.number_format = "0.0%"
    rtop = math.ceil(rmax * 100 + 0.5) / 100
    ch2.y_axis.scaling.min = 0; ch2.y_axis.scaling.max = rtop; ch2.y_axis.majorUnit = rtop / 6
    for s, col, nm in zip(ch2.series, ["2E75B6", "548235"], ["한국 기준금리(보조축)", "미국 기준금리 상단(보조축)"]):
        s.tx = SeriesLabel(v=nm); style_line(s, col, 28575, "dash", "diamond")
    ch.legend.position = "b"; secondary(ch, ch2); ws.add_chart(ch, "A4")

    ch3 = LineChart(); ch3.title = "② 아파트 매매 거래량 증감률(전반기 대비) vs 한·미 기준금리"; ch3.height = 15; ch3.width = 36
    ch3.y_axis.title = "거래량 증감률 (전반기 대비)"; ch3.y_axis.number_format = "0%"; ch3.y_axis.majorGridlines = None
    hn = [v["n"] for v in M["_halves"].values()]
    chg = [(hn[i] - hn[i - 1]) / max(hn[i - 1], 1) for i in range(1, len(hn))] or [1]
    vmax = max(2, math.ceil(max(chg))); ch3.y_axis.scaling.min = -1; ch3.y_axis.scaling.max = vmax; ch3.y_axis.majorUnit = 1
    ch3.x_axis.title = "반기말 시점 (6개월 단위)"; ch3.x_axis.tickLblPos = "low"
    for col in (6, 7, 8): ch3.add_data(Reference(wv, min_col=col, min_row=VS + 1, max_row=VL), titles_from_data=False)
    ch3.set_categories(Reference(wv, min_col=1, min_row=VS + 1, max_row=VL))
    for s, col, nm in zip(ch3.series, ["C00000", "ED7D31", "7030A0"], [f"{HN} 거래량 증감률(추정)", f"{LN} 거래량 증감률(추정)", f"{NAME} 거래량 증감률(실거래)"]):
        s.tx = SeriesLabel(v=nm); style_line(s, col, 31750, marker="triangle")
    ch4 = LineChart(); ch4.add_data(Reference(wv, min_col=9, max_col=10, min_row=VS + 1, max_row=VL), titles_from_data=False)
    ch4.y_axis.title = "기준금리 (보조축)"; ch4.y_axis.number_format = "0.0%"
    ch4.y_axis.scaling.min = -(rtop / vmax); ch4.y_axis.scaling.max = rtop
    for s, col, nm in zip(ch4.series, ["2E75B6", "548235"], ["한국 기준금리(보조축)", "미국 기준금리 상단(보조축)"]):
        s.tx = SeriesLabel(v=nm); style_line(s, col, 28575, "dash", "diamond")
    ch3.legend.position = "b"; secondary(ch3, ch4)
    ws["A37"] = "② 거래량 증감률 (위 차트와 X축 동일, 엑셀 차트는 세로축이 2개까지라 별도 패널로 배치)"; ws["A37"].font = font(bold=True, size=11, color=NAVY)
    ws.add_chart(ch3, "A38")
    r = 70; ws[f"A{r}"] = "핵심 분석"; ws[f"A{r}"].font = font(bold=True, size=13, color=NAVY); r += 1
    for t, b in TXT.get("chart_notes", []):
        ws[f"A{r}"] = t; ws[f"A{r}"].font = font(bold=True, size=11)
        ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=18)
        ws[f"A{r+1}"] = b; ws[f"A{r+1}"].font = font(size=10); ws[f"A{r+1}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r + 1].height = 36; r += 2
    ws[f"A{r+1}"] = f"※ 금리·{NAME} 실거래(평단가·거래량)는 실제 값, 구 그룹 평단가·거래량은 공표치 기반 추정치입니다."
    ws[f"A{r+1}"].font = font(size=9, italic=True, color="C00000")
    for c in "ABCDEFGHIJKLMNOPQR": ws.column_dimensions[c].width = 10.5
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0

    # ===================== 출처_가정 =====================
    wn["A1"] = "출처 및 가정"; wn["A1"].font = font(bold=True, size=14)
    rows = [("구분", "내용")] + [tuple(x) for x in mkt.get("sources_rows", [])] + policy_rows() + [tuple(x) for x in fill(apt.get("sources_rows", []), M)]
    for i, (a, b) in enumerate(rows):
        rr = 3 + i; wn[f"A{rr}"] = a; wn[f"B{rr}"] = b
        for c in "AB": wn[f"{c}{rr}"].border = box; wn[f"{c}{rr}"].alignment = Alignment(wrap_text=True, vertical="top")
        if i == 0:
            for c in "AB": wn[f"{c}{rr}"].font = font(bold=True, color="FFFFFF"); wn[f"{c}{rr}"].fill = HDR
        else:
            wn[f"A{rr}"].font = font(bold=True); wn[f"B{rr}"].font = font(); wn.row_dimensions[rr].height = 48
    wn.column_dimensions["A"].width = 26; wn.column_dimensions["B"].width = 115

    # ===================== 단지 분석 시트 =====================
    ctx = dict(wb=wb, wd=wd, wv=wv, apt=apt, mkt=mkt, TXT=TXT, M=M, trades=trades, SN_T=SN_T, SN_A=SN_A, TS=TS, TE=TE,
               TYROW=TYROW, WROW=WROW, S=S, L=L, VS=VS, VL=VL, periods=periods, vol=vol, rng=R)
    build_analysis_sheet(ctx)
    wb.active = wb.sheetnames.index(SN_A)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path, M


# ======================================================================
def build_analysis_sheet(C):
    wb, wd, wv, apt, mkt, TXT, M = C["wb"], C["wd"], C["wv"], C["apt"], C["mkt"], C["TXT"], C["M"]
    SN, SN_T, TS, TE, S, L, VS, VL = C["SN_A"], C["SN_T"], C["TS"], C["TE"], C["S"], C["L"], C["VS"], C["VL"]
    trades, periods, vol, TYROW, WROW, R = C["trades"], C["periods"], C["vol"], C["TYROW"], C["WROW"], C["rng"]
    NAME = apt["name"]; reg = apt.get("region", {}); types = {t["pyeong"]: t for t in apt["types"]}
    focus = apt["user_deal"]["pyeong"]; comp = apt.get("compare_pyeong"); ft = types[focus]
    G = mkt["groups"]
    cs = wb.create_sheet(SN, 1)
    thin2 = Side(style="thin", color="D0D0D0"); bx = Border(left=thin2, right=thin2, top=thin2, bottom=thin2)
    cs.column_dimensions["A"].width = 2
    for c in "BCDEFGHIJKLM": cs.column_dimensions[c].width = 11.5
    cs.column_dimensions["N"].width = 2; cs.sheet_view.showGridLines = False
    ROW = [1]; B, Mx = 2, 13

    def nxt(n=1):
        r = ROW[0]; ROW[0] += n; return r

    def lines(text, span):
        cpl = max(8, int(span * 11.5 / 1.95)); n = 0
        for part in str(text).split("\n"): n += max(1, math.ceil(len(part) / cpl))
        return n

    def mw(r, c1, c2, val, f=None, fill=None, al=None, border=False, fmt=None):
        if c2 > c1: cs.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        cell = cs.cell(r, c1, val); cell.font = f or font(size=10)
        cell.alignment = al or Alignment(wrap_text=True, vertical="center")
        if fill:
            for cc in range(c1, c2 + 1): cs.cell(r, cc).fill = FILL(fill)
        if border:
            for cc in range(c1, c2 + 1): cs.cell(r, cc).border = bx
        if fmt: cell.number_format = fmt
        return cell

    def banner(num, text):
        nxt(); r = nxt()
        mw(r, B, Mx, f"{num}  {text}", f=font(bold=True, size=14, color="FFFFFF"), fill=NAVY, al=Alignment(vertical="center", indent=1))
        cs.row_dimensions[r].height = 30; return r

    def sub(text):
        nxt(); r = nxt()
        mw(r, B, Mx, text, f=font(bold=True, size=11.5, color=NAVY), fill=LBLUE, al=Alignment(vertical="center", indent=1))
        cs.row_dimensions[r].height = 22; return r

    def callout(text, fill=CREAM):
        if not text: return None
        r = nxt()
        mw(r, B, Mx, text, f=font(bold=True, size=10.5), fill=fill, al=Alignment(wrap_text=True, vertical="center", indent=1))
        cs.cell(r, B).border = Border(left=Side(style="thick", color=ORANGE))
        cs.row_dimensions[r].height = lines(text, 12) * 15 + 10; return r

    def bullet(text, mark="▪", color="000000"):
        r = nxt(); t = f"{mark} {text}"
        mw(r, B, Mx, t, f=font(size=10, color=color), al=Alignment(wrap_text=True, vertical="top", indent=1))
        cs.row_dimensions[r].height = lines(t, 12) * 14.5 + 5; return r

    def note(text):
        r = nxt(); mw(r, B, Mx, text, f=font(size=9, italic=True, color="7F7F7F"), al=Alignment(wrap_text=True, vertical="top"))
        cs.row_dimensions[r].height = lines(text, 12) * 13 + 4; return r

    def table(spec, rows, hdr_fill=BLUE2, zebra=True, first_bold=True, min_h=20, fmts=None, fills=None):
        r = nxt(); c = B; hm = 1
        for span, h in spec:
            mw(r, c, c + span - 1, h, f=font(bold=True, size=10, color="FFFFFF"), fill=hdr_fill,
               al=Alignment(horizontal="center", vertical="center", wrap_text=True), border=True)
            hm = max(hm, lines(h, span)); c += span
        cs.row_dimensions[r].height = max(22, hm * 14 + 8)
        out = []
        for i, row in enumerate(rows):
            r = nxt(); c = B; hm = 1
            for j, ((span, _), v) in enumerate(zip(spec, row)):
                fl = (fills[i][j] if fills and fills[i] and fills[i][j] else (GREY if zebra and i % 2 == 1 else None))
                isnum = not (isinstance(v, str) and not v.startswith("="))
                al = Alignment(horizontal=("center" if (isnum or span == 1) else "left"), vertical="center", wrap_text=True, indent=0 if isnum else 1)
                fnt = font(size=10, bold=(first_bold and j == 0))
                if isinstance(v, str) and v.startswith("=") and "!" in v: fnt = Font(name=F, size=10, color="008000", bold=(first_bold and j == 0))
                mw(r, c, c + span - 1, v, f=fnt, fill=fl, al=al, border=True, fmt=(fmts[j] if fmts and j < len(fmts) and fmts[j] else None))
                if not (isinstance(v, str) and v.startswith("=")): hm = max(hm, lines(v, span))
                c += span
            cs.row_dimensions[r].height = max(min_h, hm * 14 + 6); out.append(r)
        return out

    def inputs(rows_):
        addr = {}
        for key, lab, val, fm, nt in rows_:
            r = nxt()
            mw(r, B, 5, lab, f=font(bold=True), fill=PALE, border=True, al=Alignment(vertical="center", indent=1))
            isf = isinstance(val, str) and val.startswith("=")
            mw(r, 6, 7, val, f=(font(bold=True) if isf else Font(name=F, size=10.5, bold=True, color="0000FF")),
               fill=(None if isf else "FFFF99"), border=True, fmt=fm, al=Alignment(horizontal="center", vertical="center"))
            mw(r, 8, Mx, nt, f=font(size=9, color="595959"), border=True, al=Alignment(indent=1, vertical="center", wrap_text=True))
            cs.row_dimensions[r].height = 20; addr[key] = f"$F${r}"
        return addr

    # ---------------- 표지
    r = nxt(); mw(r, B, Mx, f"{NAME} 종합 분석 리포트  |  단지·{reg.get('dong','지역')}·금리·가격 시나리오",
                  f=font(bold=True, size=18, color="FFFFFF"), fill=NAVY, al=Alignment(vertical="center", indent=1)); cs.row_dimensions[r].height = 40
    r = nxt(); mw(r, B, Mx, f"{apt.get('address','')} · {apt.get('households','')}세대 · {apt.get('built','')} 입주  |  작성 기준일 {mkt['as_of']}  |  초록 글씨 = 수식 연결, 노란 칸 = 사용자 입력  |  투자 권유가 아닌 정보 제공용 분석",
                  f=font(size=9, color="595959"), fill=GREY, al=Alignment(vertical="center", indent=1)); cs.row_dimensions[r].height = 20
    nxt(); r = nxt(); mw(r, B, Mx, "목차 (클릭하면 이동)", f=font(bold=True, size=11, color=NAVY), al=Alignment(indent=1))
    SECTIONS = [("summary", "핵심 요약 대시보드"), ("price", f"사용자 거래({M['user_price_short']}) 가격 위치 진단"), ("complex", "단지 전체 분석"),
                ("episode", (TXT.get("episode") or {}).get("title", "")), ("spread", "한·미 금리차와 거래량의 관계"),
                ("region", (TXT.get("region") or {}).get("title", "")), ("outlook", f"{M['user_price_short']} 거래 기준 가격 전망 & 미래 시나리오"), ("sources", "출처 및 유의사항")]
    SECTIONS = [(k, t) for k, t in SECTIONS if not (k in ("episode", "region") and not TXT.get(k))]
    ROMAN = ["Ⅰ", "Ⅱ", "Ⅲ", "Ⅳ", "Ⅴ", "Ⅵ", "Ⅶ", "Ⅷ", "Ⅸ"]
    NUM = {k: ROMAN[i] for i, (k, _) in enumerate(SECTIONS)}; TITLE = dict(SECTIONS)
    TOC = [nxt() for _ in SECTIONS]; SEC = {}

    # ---------------- Ⅰ 요약
    SEC["summary"] = banner(NUM["summary"] + ".", TITLE["summary"]); nxt(); KPI_R = ROW[0]; nxt(5)
    sub("한눈에 보는 결론")
    for s_ in TXT.get("summary", []): bullet(s_)

    # ---------------- Ⅱ 가격 진단
    SEC["price"] = banner(NUM["price"] + ".", TITLE["price"]); sub("입력값")
    pr = apt.get("price_refs", {}); nb = pr.get("new_build"); pp = pr.get("private_price")
    ud = apt["user_deal"]
    rows_in = [("P0", "거래가격 (만원)", ud["price"], "#,##0", ud.get("note", "사용자 입력")),
               ("UNIT", "물건 (동·호·층)", M.get("user_unit", ""), None, f"{M.get('user_side','미확인')} 검토" + (f" · {ud['unit_note']}" if ud.get("unit_note") else "")),
               ("PY", "공급평형 (평)", focus, "0", f"{focus}평형"),
               ("AREA", "전용면적 (㎡)", ft["area"], "0.0", f"전용 약 {ft['area']/PY_M2:.1f}평")]
    if pp: rows_in.append(("PRIV", pp["label"] + " (만원)", pp["value"], "#,##0", pp.get("source", "")))
    if ft.get("kb_jeonse"): rows_in.append(("JEON", f"{focus}평 전세 KB시세 (만원, {apt.get('kb_asof','')})", ft["kb_jeonse"], "#,##0", ft.get("jeonse_note", "")))
    if nb:
        rows_in.append(("NB", nb["label"] + " (만원)", nb["value"], "#,##0", nb.get("source", "")))
        if nb.get("member_ratio"): rows_in.append(("NBR", "조합원분양가 ÷ 일반분양가 비율", nb["member_ratio"], "0%", nb.get("ratio_source", "")))
    IN = inputs(rows_in); P0, AREA = IN["P0"], IN["AREA"]
    sub("비교 기준별 위치")
    spec = [(4, "비교 기준"), (2, "가격 (만원)"), (2, "거래가가 기준보다"), (2, "전용 평단가 (만원)"), (2, "비고")]
    start = ROW[0] + 1; rows = []; fills = []; tag = {}
    def addrow(k, lab, price_f, ppp_f, nt, hl=False):
        rr = start + len(rows); tag[k] = rr
        rows.append([lab, price_f, (f"={P0}/F{rr}-1" if k != "user" else "=0"), ppp_f.replace("{r}", str(rr)), nt])
        fills.append(["FFFF99"] * 5 if hl else None)
    mx = M.get("_focus_max")
    addrow("max", f"{focus}평 실거래 최고가 (국토부)", f"=_xlfn.MAXIFS({R('G')},{R('D')},{focus},{R('I')},1)", f"=F{{r}}/({AREA}/3.3058)",
           f"{mx['date']:%Y.%m.%d} · {mx['floor']}층" if mx and mx["floor"] else "")
    fh = sorted({t["half"] for t in trades if t["pyeong"] == focus and t["include"]})
    if fh:
        lh = fh[-1]; nlh = sum(1 for t in trades if t["pyeong"] == focus and t["include"] and t["half"] == lh)
        addrow("havg", f"{focus}평 {lh} 반기 평균 실거래", f'=AVERAGEIFS({R("G")},{R("D")},{focus},{R("B")},"{lh}",{R("I")},1)', f"=F{{r}}/({AREA}/3.3058)", f"{nlh}건 평균")
    udong = M.get("_user_dong"); dong_note = None
    if udong:
        dm = M.get("_dong_max"); dn = M.get("dong_n", "0")
        if dm:
            addrow("dmax", f"{udong}동 {focus}평 실거래 최고가", f'=_xlfn.MAXIFS({R("G")},{R("D")},{focus},{R("I")},1,{R("C")},"{udong}")', f"=F{{r}}/({AREA}/3.3058)",
                   f"{dm['date']:%Y.%m.%d} · {dm['floor']}층 · 같은 동 {dn}건" if dm["floor"] else f"같은 동 {dn}건")
            addrow("davg", f"{udong}동 {focus}평 평균 실거래", f'=AVERAGEIFS({R("G")},{R("D")},{focus},{R("I")},1,{R("C")},"{udong}")', f"=F{{r}}/({AREA}/3.3058)", f"동 표기 거래 {dn}건 평균(전 기간)")
        else:
            dong_note = f"※ {udong}동 {focus}평은 국토부 자료에 동 표기 거래가 없어 단지 전체 기준으로 판단합니다."
    if ft.get("kb_price"): addrow("kb", f"{focus}평 KB 매매시세 ({apt.get('kb_asof','')})", f"='{SN_T}'!$P${TYROW[focus]}", f"=F{{r}}/({AREA}/3.3058)", "KB 일반평균가")
    if pp: addrow("priv", pp["label"], f"={IN['PRIV']}", f"=F{{r}}/({AREA}/3.3058)", "민간 시세")
    addrow("user", "사용자 거래가", f"={P0}", f"=F{{r}}/({AREA}/3.3058)", "검토 대상", hl=True)
    if comp and comp in types:
        ca = types[comp]["area"]
        addrow("cmax", f"{comp}평(전용 {ca}㎡) 실거래 최고가", f"=_xlfn.MAXIFS({R('G')},{R('D')},{comp},{R('I')},1)", f"=F{{r}}/({ca}/3.3058)", "참고")
        addrow("cconv", f"{comp}평을 사용자 거래 평단가로 환산", f"=J{tag['user']}*({ca}/3.3058)", f"=J{tag['user']}", "소형 평형 프리미엄 크기 확인용")
    if nb:
        a = nb.get("area", 59)
        addrow("nb", nb["label"], f"={IN['NB']}", f"=F{{r}}/({a}/3.3058)", nb.get("note", ""))
        if nb.get("member_ratio"):
            addrow("nbg", nb.get("general_label", "일반분양가 환산"), f"={IN['NB']}/{IN['NBR']}", f"=F{{r}}/({a}/3.3058)", "조합원가 ÷ 비율 역산")
    addrow("low", f"{G['low']['name']} 평균 평단가 × 전용{ft['area']/PY_M2:.1f}평", f"=데이터!$D${L}*({AREA}/3.3058)", f"=데이터!$D${L}", "'데이터' 시트 추정치")
    addrow("high", f"{G['high']['name']} 평균 평단가 × 전용{ft['area']/PY_M2:.1f}평", f"=데이터!$C${L}*({AREA}/3.3058)", f"=데이터!$C${L}", "'데이터' 시트 추정치")
    table(spec, rows, fmts=[None, "#,##0", '+0.0%;-0.0%;0.0%', "#,##0", None], fills=fills)
    if dong_note: note(dong_note)
    JR = None
    if "JEON" in IN:
        JR = nxt()
        mw(JR, B, 5, "전세가율 (KB 전세 ÷ 거래가)", f=font(bold=True), fill=PALE, border=True, al=Alignment(indent=1, vertical="center"))
        mw(JR, 6, 7, f"={IN['JEON']}/{P0}", f=font(bold=True, size=11), border=True, al=Alignment(horizontal="center"), fmt="0.0%")
        mw(JR, 8, Mx, "전세가율이 낮을수록 하락기 하방 지지력이 약함", f=font(size=9, color="595959"), border=True, al=Alignment(indent=1, vertical="center"))
    nxt(); callout(TXT.get("price_callout", ""))

    # ---------------- Ⅲ 단지 전체
    cx = TXT.get("complex", {})
    SEC["complex"] = banner(NUM["complex"] + ".", f"단지 전체 분석 ({trades[0]['date'].year}~{parse_date(mkt['as_of']).year})")
    sub("1. 단지 개요")
    if cx.get("overview"): table([(3, "항목"), (9, "내용")], cx["overview"], min_h=22)
    if cx.get("overview_source"): note(cx["overview_source"])
    sub("2. 연도별 실거래 요약 (국토부, 이상거래 제외 · 수식 연결)")
    ct = comp if comp in types else None
    spec = [(1, "연도"), (1, "총 거래"), (1, f"{focus}평 건수"), (2, f"{focus}평 평균 (만원)"), (1, f"{focus}평 최고"), (1, f"{focus}평 최저")]
    spec += [(1, f"{ct}평 건수"), (2, f"{ct}평 평균 (만원)"), (2, f"{ct}평 최고")] if ct else [(5, "비고")]
    Y0 = ROW[0] + 1; rows = []
    for k, y in enumerate(range(trades[0]["date"].year, parse_date(mkt["as_of"]).year + 1)):
        rr = Y0 + k; lo_ = f'">="&DATE({y},1,1)'; hi_ = f'"<"&DATE({y+1},1,1)'
        base_ = f"{R('A')},{lo_},{R('A')},{hi_}"
        row = [str(y), f"=COUNTIFS({base_})", f"=COUNTIFS({base_},{R('D')},{focus},{R('I')},1)",
               f'=IFERROR(AVERAGEIFS({R("G")},{base_},{R("D")},{focus},{R("I")},1),"-")',
               f'=IF(D{rr}=0,"-",_xlfn.MAXIFS({R("G")},{base_},{R("D")},{focus},{R("I")},1))',
               f'=IF(D{rr}=0,"-",_xlfn.MINIFS({R("G")},{base_},{R("D")},{focus},{R("I")},1))']
        if ct:
            row += [f"=COUNTIFS({base_},{R('D')},{ct},{R('I')},1)", f'=IFERROR(AVERAGEIFS({R("G")},{base_},{R("D")},{ct},{R("I")},1),"-")',
                    f'=IF(I{rr}=0,"-",_xlfn.MAXIFS({R("G")},{base_},{R("D")},{ct},{R("I")},1))']
        else:
            row += [""]
        rows.append(row)
    table(spec, rows, fmts=[None, "0", "0", "#,##0", "#,##0", "#,##0", "0", "#,##0", "#,##0"])
    note("※ 마지막 해는 as_of까지 신고분. 해당 평형 거래가 없는 해는 '-'. 총 거래는 이상거래 포함 전체 건수.")
    sub("3. 차트: 단지 평균 평단가와 반기 거래건수")
    CH_R = nxt(); nxt(22)
    bar = BarChart(); bar.type = "col"
    bar.add_data(Reference(wv, min_col=5, min_row=VS + 1, max_row=VL - 1), titles_from_data=False)
    bar.series[0].tx = SeriesLabel(v="반기 거래건수(보조축)"); bar.series[0].graphicalProperties.solidFill = "BDD7EE"; bar.series[0].graphicalProperties.line.solidFill = "BDD7EE"
    bar.y_axis.title = "거래건수 (건)"
    ln = LineChart(); ln.add_data(Reference(wd, min_col=7, min_row=S, max_row=L - 1), titles_from_data=False)
    style_line(ln.series[0], "7030A0", 34925); ln.series[0].tx = SeriesLabel(v=f"{NAME} 평균 평단가(만원/3.3㎡, 전용)")
    hmax_ppp = max([h["avg_ppp"] or 0 for h in M["_halves"].values()] + [1000])
    ln.y_axis.title = "평단가 (만원)"; ln.y_axis.scaling.min = 0; ln.y_axis.scaling.max = nice_max(hmax_ppp, 1000); ln.y_axis.number_format = "#,##0"; ln.y_axis.majorGridlines = None
    nmax = max([h["n"] for h in M["_halves"].values()] + [5]); bar.y_axis.scaling.min = 0; bar.y_axis.scaling.max = nice_max(nmax, 5)
    ln.set_categories(Reference(wd, min_col=1, min_row=S, max_row=L - 1)); bar.set_categories(Reference(wd, min_col=1, min_row=S, max_row=L - 1))
    ln.title = "평균 평단가 vs 반기 거래건수"; ln.height = 11; ln.width = 25; ln.x_axis.tickLblPos = "low"; ln.legend.position = "b"
    secondary(ln, bar); cs.add_chart(ln, f"B{CH_R}")
    if cx.get("phases"):
        sub("4. 사이클별 흐름 해석")
        table([(2, "국면"), (2, "금리·정책 환경"), (4, "단지 가격·거래"), (4, "해석")], cx["phases"], min_h=36)
    if cx.get("swot"):
        sub("5. 입지·상품 SWOT")
        sw = cx["swot"]
        table([(6, "S 강점"), (6, "W 약점")], [[sw.get("S", ""), sw.get("W", "")]], hdr_fill="548235", zebra=False, first_bold=False, min_h=60)
        table([(6, "O 기회"), (6, "T 위협")], [[sw.get("O", ""), sw.get("T", "")]], hdr_fill="C00000", zebra=False, first_bold=False, min_h=60)

    # ---------------- Ⅳ 에피소드(선택)
    ep = TXT.get("episode")
    if ep:
        SEC["episode"] = banner(NUM["episode"] + ".", ep["title"]); callout(ep.get("callout", ""))
        keys = ep.get("fact_periods", [])[:4]
        if keys:
            sub("1. 당시 팩트 (수식 연결)")
            labs = [p[0] for p in periods]; vlabs = [v["label"] for v in vol]
            spec = [(4, "구분")] + [(2, k) for k in keys]
            defs = [("한국 기준금리", "I", "0.00%"), ("미국 기준금리 상단", "J", "0.00%"), ("한·미 금리차 (%p)", "K", '+0.00;-0.00;0.00'), ("단지 거래건수", "E", "0")]
            rows = [[nm] + [f"=데이터!{col}{S+labs.index(k)}" for k in keys] for nm, col, _ in defs]
            rows.append(["거래량 증감률 (전반기 대비)"] + [f"=거래량!H{VS+vlabs.index(k)}" for k in keys])
            outs = table(spec, rows)
            for rr, fm in zip(outs, [d[2] for d in defs] + ['+0%;-0%;0%']):
                for cc in range(6, 6 + 2 * len(keys), 2): cs.cell(rr, cc).number_format = fm
        y = ep.get("year")
        if y:
            sub(f"2. {y}년 실거래와 당시 이벤트")
            idx = [i for i, t in enumerate(trades) if t["date"].year == y]
            spec = [(2, "계약일"), (1, "평형"), (1, "층"), (2, "거래가 (만원)"), (2, f"{y}년 이전 동일평형 최고가 대비"), (4, "당시 이벤트")]
            rows = []
            for i in idx:
                tr = TS + i; t = trades[i]
                peak = f"_xlfn.MAXIFS({R('G')},{R('D')},'{SN_T}'!D{tr},{R('A')},\"<\"&DATE({y},1,1),{R('I')},1)"
                rows.append([f"='{SN_T}'!A{tr}", f"='{SN_T}'!D{tr}", f"='{SN_T}'!F{tr}", f"='{SN_T}'!G{tr}",
                             f"=IFERROR('{SN_T}'!G{tr}/{peak}-1,\"-\")", ep.get("trade_events", {}).get(t["date"].isoformat(), "")])
            if rows: table(spec, rows, fmts=["yyyy-mm-dd", "0", "0", "#,##0", '+0%;-0%;0%', None])
        if ep.get("factors"):
            sub("3. 원인 분석 (당시 기사 기반)")
            table([(2, "요인"), (1, "시기"), (5, "기사 요지 (출처)"), (4, f"{NAME} 함의")], ep["factors"], min_h=40)
        if ep.get("insights"):
            sub("4. 전문가 관점 해석")
            for t_ in ep["insights"]: bullet(t_)

    # ---------------- Ⅴ 금리차
    sp = TXT.get("spread", {})
    SEC["spread"] = banner(NUM["spread"] + ".", TITLE["spread"]); callout(sp.get("callout", ""))
    sub("1. 상관계수 (반기 관측치, 엑셀 CORREL)")
    HC = 16; hr = ROW[0]; cs.cell(hr, HC, "보조 데이터(상관분석용)").font = font(size=8, color="7F7F7F")
    for i, h in enumerate(["반기", "한국", "미국", "금리차", "Δ금리차", "Δ한국", "단지", "고가", "저가", "Δ단지"]):
        cs.cell(hr + 1, HC + i, h).font = font(size=8, bold=True, color="7F7F7F")
    n = len([p for p in periods if not p[2]]); D0 = hr + 2; DL = D0 + n - 1
    for i in range(n):
        rr = D0 + i; drow = S + i; vrow = VS + 1 + i
        vals = [f"=데이터!A{drow}", f"=데이터!I{drow}*100", f"=데이터!J{drow}*100", f"=Q{rr}-R{rr}",
                (f"=S{rr}-S{rr-1}" if i else None), (f"=Q{rr}-Q{rr-1}" if i else None),
                f"=거래량!E{vrow}", f"=거래량!C{vrow}", f"=거래량!D{vrow}", (f"=V{rr}-V{rr-1}" if i else None)]
        for j, v in enumerate(vals):
            if v is not None: cs.cell(rr, HC + j, v).font = font(size=8, color="7F7F7F")
    for j, w in enumerate([8, 6, 6, 6, 6, 6, 6, 7, 7, 6]): cs.column_dimensions[get_column_letter(HC + j)].width = w
    Rr = lambda col, a=D0, b=DL: f"${col}${a}:${col}${b}"
    spec = [(4, "설명변수 ＼ 거래량"), (2, f"{NAME} (실거래)"), (2, f"{G['high']['name']} (추정)"), (2, f"{G['low']['name']} (추정)"), (2, "읽는 법")]
    rows = [["한·미 금리차 (같은 반기)"] + [f"=CORREL({Rr('S')},{Rr(c)})" for c in "VWX"] + ["0에 가까우면 관계 약함"],
            ["한·미 금리차 (6개월 선행)"] + [f"=CORREL($S${D0}:$S${DL-1},${c}${D0+1}:${c}${DL})" for c in "VWX"] + ["선행성 확인"],
            ["한국 기준금리 수준"] + [f"=CORREL({Rr('Q')},{Rr(c)})" for c in "VWX"] + ["음(-)이면 금리↑ 거래↓"],
            ["미국 기준금리 수준"] + [f"=CORREL({Rr('R')},{Rr(c)})" for c in "VWX"] + ["간접 효과"],
            ["Δ금리차 vs Δ거래량", f"=CORREL({Rr('T',D0+1)},{Rr('Y',D0+1)})", "-", "-", "변화분끼리 비교"],
            ["Δ한국금리 vs Δ거래량", f"=CORREL({Rr('U',D0+1)},{Rr('Y',D0+1)})", "-", "-", "금리 변화 방향 효과"]]
    table(spec, rows, fmts=[None] + ['+0.00;-0.00;0.00'] * 3 + [None])
    crit_r = 2 / math.sqrt(n) if n else 0.47
    note(f"※ 관측치 {n}개에서 5% 유의수준 기준 |r| ≥ 약 {min(0.99, 1.96/math.sqrt(max(n-1,1))):.2f} 필요. 구 그룹 거래량은 추정치라 참고용. 보조 데이터는 이 시트 P~Y열.")
    sub("2. 금리차 구간별 평균 거래량 (반기당)")
    regs = [("한국 > 미국 (+)", '">0"', None), ("0 ~ -0.75%p", '"<=0"', '">-1"'), ("-1%p 이하 (역전 심화)", '"<=-1"', None)]
    rows = []
    for nm, c1, c2 in regs:
        crit = f"{Rr('S')},{c1}" + (f",{Rr('S')},{c2}" if c2 else "")
        rows.append([nm, f"=COUNTIFS({crit})"] + [f"=IFERROR(AVERAGEIFS({Rr(c)},{crit}),\"-\")" for c in "VWX"])
    table([(3, "금리차 구간"), (1, "반기 수"), (3, NAME), (3, f"{G['high']['name']} (추정)"), (2, f"{G['low']['name']} (추정)")], rows,
          fmts=[None, "0", "0.0", "#,##0", "#,##0"])
    if sp.get("interpretation"):
        sub("3. 해석")
        for t_ in sp["interpretation"]: bullet(t_)

    # ---------------- Ⅵ 지역(선택)
    rg = TXT.get("region")
    if rg:
        SEC["region"] = banner(NUM["region"] + ".", rg["title"]); callout(rg.get("callout", ""))
        k = 1
        if rg.get("timeline"):
            sub(f"{k}. 정책·개발 변천사"); k += 1; table([(2, "시기"), (10, "내용")], rg["timeline"])
        if rg.get("pipeline"):
            sub(f"{k}. 개발·정비사업 파이프라인"); k += 1
            table([(3, "사업"), (2, "규모"), (3, "진행 단계"), (4, f"{NAME}와의 관계")], rg["pipeline"], min_h=30)
            if rg.get("pipeline_source"): note(rg["pipeline_source"])
        if rg.get("structure"):
            sub(f"{k}. 수요·공급 구조 진단"); k += 1
            for t_ in rg["structure"]: bullet(t_)
        if rg.get("outlook"):
            sub(f"{k}. 지역 전망 (시계열)"); k += 1
            table([(2, "기간"), (2, "국면"), (4, "예상 흐름"), (4, f"{NAME} 영향")], rg["outlook"], min_h=48)

    # ---------------- Ⅶ 전망·시나리오
    ol = TXT.get("outlook", {})
    SEC["outlook"] = banner(NUM["outlook"] + ".", TITLE["outlook"])
    if ol.get("drivers"):
        sub(f"1. 가격 드라이버 점검표 ({mkt['as_of'][:7]} 기준)")
        dfl = [[None, None, c_, c_, None] for c_ in [(GREENF if d[2] == "↑" else REDF if d[2] == "↓" else None) for d in ol["drivers"]]]
        table([(2, "요인"), (5, "현재 상황"), (1, "방향"), (1, "강도"), (3, "근거·출처")], ol["drivers"], min_h=36, fills=dfl)
        callout(ol.get("drivers_callout", ""))
    sub("2. 시나리오 (확률·변동률은 분석자 주관 가정 → 노란 칸 수정 가능)")
    spec = [(2, "시나리오"), (1, "확률"), (1, "1년 변동"), (1, "3년 누적"), (2, "1년 후 가격 (만원)"), (2, "3년 후 가격 (만원)"), (3, "전제·트리거")]
    S0 = ROW[0] + 1; rows = []; fl = []
    for i, (nm, p, y1, y3, tx) in enumerate(ol.get("scenarios", [])):
        rr = S0 + i
        rows.append([nm, p, y1, y3, f"={P0}*(1+E{rr})", f"={P0}*(1+F{rr})", tx]); fl.append([None, "FFFF99", "FFFF99", "FFFF99", None, None, None])
    outs = table(spec, rows, fmts=[None, "0%", '+0%;-0%;0%', '+0%;-0%;0%', "#,##0", "#,##0", None], fills=fl, min_h=40)
    for rr in outs:
        for cc in (4, 5, 6): cs.cell(rr, cc).font = Font(name=F, size=10, bold=True, color="0000FF")
    EXP = nxt()
    if outs:
        a, b_ = outs[0], outs[-1]
        mw(EXP, B, 3, "확률가중 기대값", f=font(bold=True, color="FFFFFF"), fill=NAVY, border=True, al=Alignment(indent=1, vertical="center"))
        mw(EXP, 4, 4, f"=SUM(D{a}:D{b_})", f=font(bold=True), border=True, fmt="0%", al=Alignment(horizontal="center"))
        mw(EXP, 5, 5, f"=SUMPRODUCT($D${a}:$D${b_},E{a}:E{b_})", f=font(bold=True), border=True, fmt='+0.0%;-0.0%;0.0%', al=Alignment(horizontal="center"))
        mw(EXP, 6, 6, f"=SUMPRODUCT($D${a}:$D${b_},F{a}:F{b_})", f=font(bold=True), border=True, fmt='+0.0%;-0.0%;0.0%', al=Alignment(horizontal="center"))
        mw(EXP, 7, 8, f"={P0}*(1+E{EXP})", f=font(bold=True), border=True, fmt="#,##0", al=Alignment(horizontal="center"))
        mw(EXP, 9, 10, f"={P0}*(1+F{EXP})", f=font(bold=True), border=True, fmt="#,##0", al=Alignment(horizontal="center"))
        mw(EXP, 11, Mx, "확률 합계가 100%인지 확인", f=font(size=9, color="595959"), border=True, al=Alignment(indent=1, vertical="center"))
        r = nxt()
        mw(r, B, 7, "참고: 거래가가 시장가로 확인되지 않고 직전 최고가 수준으로 평가될 경우 즉시 평가손", f=font(size=10), fill=REDF, border=True, al=Alignment(indent=1, vertical="center"))
        mw(r, 8, 9, f"=F{tag['max']}/{P0}-1", f=font(bold=True, color="C00000"), fill=REDF, border=True, fmt='+0.0%;-0.0%;0.0%', al=Alignment(horizontal="center"))
        mw(r, 10, Mx, "KB시세 기준 평가손은 Ⅱ장 표 참고", f=font(size=9, color="595959"), fill=REDF, border=True, al=Alignment(indent=1, vertical="center"))
    md = ol.get("model", {})
    pol = load_policy(); PT, PB, PL = pol["acquisition_tax"], pol["brokerage"], pol["loan"]
    EDU = 1 + PT["local_education_tax_ratio"]; VAT = 1 + PB["vat_ratio"]; SELL = PB["sell_rate_simplified"]
    sub("3. 3년 보유 손익 모델 (1세대 1주택·실거주 가정, 노란 칸 수정 가능)")
    base_r = ROW[0]
    keys = ["LTV", "CAP", "LOAN", "RATE", "TAXR", "TAX", "BRK", "ETC", "ACQ", "CASH", "HOLD", "OPP", "SAVE"]
    MI = {k: f"$F${base_r+i}" for i, k in enumerate(keys)}; MI["P0"] = P0
    model = [("LTV", "대출 LTV", md.get("ltv", PL["ltv_default"]), "0%", md.get("ltv_note", "")),
             ("CAP", "주담대 한도 (만원)", md.get("loan_cap", PL["loan_cap_default"]), "#,##0", md.get("cap_note", "")),
             ("LOAN", "대출금액 (만원)", "=MIN({P0}*{LTV},{CAP})".format(**MI), "#,##0", "LTV와 한도 중 작은 값"),
             ("RATE", "대출금리 (연)", md.get("rate", PL["rate_default"]), "0.0%", md.get("rate_note", "")),
             ("TAXR", "취득세율 (지방교육세 제외)", "=" + acquisition_tax_formula(P0, pol), "0.00%", f"{PT['extra_note']} ({PT['source']}, {PT['effective_from']}~)"),
             ("TAX", "취득세+지방교육세 (만원)", "={P0}*{TAXR}*{EDU}".format(EDU=EDU, **MI), "#,##0", f"지방교육세 = 취득세의 {PT['local_education_tax_ratio']:.0%}"),
             ("BRK", "매수 중개보수 (만원)", "=" + brokerage_formula(P0, pol) + f"*{VAT}*{P0}", "#,##0", f"상한요율 + 부가세 ({PB['source']}, {PB['effective_from']}~)"),
             ("ETC", "법무·등기·기타 (만원)", md.get("etc_cost", PL["etc_cost_default"]), "#,##0", "추정치"),
             ("ACQ", "취득비용 합계 (만원)", "={TAX}+{BRK}+{ETC}".format(**MI), "#,##0", ""),
             ("CASH", "투입 현금 (만원)", "={P0}-{LOAN}+{ACQ}".format(**MI), "#,##0", "거래가 − 대출 + 취득비용"),
             ("HOLD", "연간 보유세 (만원)", md.get("holding_tax", PL["holding_tax_default"]), "#,##0", md.get("holding_note", "추정치")),
             ("OPP", "예금 기회비용 금리 (연)", md.get("deposit_rate", PL["deposit_rate_default"]), "0.0%", md.get("deposit_note", "")),
             ("SAVE", "실거주 주거비 절감 (연, 만원)", md.get("housing_saving", 0), "#,##0", md.get("saving_note", "전세·월세로 살았다면 들었을 비용"))]
    inputs([(k, lab, v, fm, nt) for k, lab, v, fm, nt in model])
    nxt()
    spec = [(2, "시나리오"), (2, "3년 후 매도가"), (1, "매도비용"), (1, "3년 이자"), (1, "3년 보유세"), (1, "주거비 절감"), (2, "순손익 (만원)"), (1, "투입현금 수익률"), (1, "기회비용 차감 후")]
    rows = []; R0 = ROW[0] + 1
    for i, rs in enumerate(outs):
        rr = R0 + i; P3 = f"D{rr}"
        rows.append([f"=B{rs}", f"=I{rs}", "=" + brokerage_formula(P3, pol) + f"*{VAT}*{P3}",
                     f"={MI['LOAN']}*{MI['RATE']}*3", f"={MI['HOLD']}*3", f"={MI['SAVE']}*3",
                     f"={P3}-F{rr}-{P0}-{MI['ACQ']}-G{rr}-H{rr}+I{rr}", f"=J{rr}/{MI['CASH']}", f"=J{rr}-{MI['CASH']}*{MI['OPP']}*3"])
    table(spec, rows, fmts=[None, "#,##0", "#,##0", "#,##0", "#,##0", "#,##0", '#,##0;[Red]-#,##0', "+0.0%;[Red]-0.0%", "#,##0;[Red]-#,##0"])
    BE = nxt()
    core = f"({P0}+{MI['ACQ']}+{MI['LOAN']}*{MI['RATE']}*3+{MI['HOLD']}*3-{MI['SAVE']}*3"
    mw(BE, B, 7, "손익분기 3년 누적 상승률 (기회비용 제외 / 포함)", f=font(bold=True, color="FFFFFF"), fill=NAVY, border=True, al=Alignment(indent=1, vertical="center"))
    mw(BE, 8, 9, f"={core})/(1-{SELL})/{P0}-1", f=font(bold=True, size=11, color="C00000"), border=True, fmt='+0.0%', al=Alignment(horizontal="center"))
    mw(BE, 10, 11, f"={core}+{MI['CASH']}*{MI['OPP']}*3)/(1-{SELL})/{P0}-1", f=font(bold=True, size=11, color="C00000"), border=True, fmt='+0.0%', al=Alignment(horizontal="center"))
    mw(BE, 12, Mx, f"매도 보수 {SELL:.2%} 가정", f=font(size=9, color="595959"), border=True, al=Alignment(indent=1, vertical="center"))
    note(f"※ 이자는 원금 상환 없이 이자만 낸다고 단순화. {pol['capital_gains']['report_note']}. 세법·대출 조건은 개인별로 다르므로 세무사·금융기관 확인 필요. (정책 기준 확인일 {pol['checked_at']})")
    callout(ol.get("model_callout", ""))
    if ol.get("checklist"):
        sub(f"4. 체크리스트 ({focus}평 {M['user_price_short']})")
        for t_ in ol["checklist"]: bullet(t_, mark="☐")
    if ol.get("monitoring"):
        sub("5. 조기 경보·모니터링 지표")
        table([(3, "지표"), (3, "현재"), (3, "경고 신호 (약세 전환)"), (3, "긍정 신호 (강세 전환)")], ol["monitoring"], min_h=30)

    # ---------------- Ⅷ 출처
    SEC["sources"] = banner(NUM["sources"] + ".", TITLE["sources"])
    for t_ in TXT.get("sources", []): bullet(t_, mark="·", color="404040")
    bullet("시나리오 확률·변동률·보유세 등은 분석자 가정입니다. 본 자료는 투자 권유가 아니며, 실제 의사결정 전 세무·대출·법률 전문가 확인이 필요합니다.", mark="·", color="404040")

    # ---------------- KPI 카드
    kb_ref = f"=F{tag['kb']}" if "kb" in tag else None
    cards = [(f"거래가 ({focus}평, {M.get('user_unit','')})", f"={P0}", '#,##0"만원"', f"전용 {ft['area']}㎡ · {M.get('user_side','미확인')}"),
             (f"{focus}평 최고가 대비", f"={P0}/F{tag['max']}-1", '+0.0%;-0.0%', f"직전 최고 {M.get('focus_max_price','')}"),
             ("KB시세 대비", (f"={P0}/F{tag['kb']}-1" if "kb" in tag else "-"), '+0.0%;-0.0%', f"KB {M.get('kb_focus','-')}"),
             ("3년 손익분기 상승률", f"=H{BE}", '+0.0%', "취득·이자·보유비 기준")]
    cards2 = [("전용 평단가", f"=J{tag['user']}", '#,##0"만원"', (f"{comp}평 최고가 평단가 대비 {M.get('user_ppp_vs_compare','')}" if comp else "")),
              ("전세가율", (f"=F{JR}" if JR else "-"), "0.0%", f"KB 전세 {M.get('jeonse_focus','-')}"),
              ("한국 / 미국 기준금리", f"=TEXT(데이터!I{L},\"0.00%\")&\" / \"&TEXT(데이터!J{L},\"0.00%\")", None, f"금리차 {M['spread_now']}"),
              ("시나리오 3년 기대 변동", (f"=F{EXP}" if outs else "-"), '+0.0%;-0.0%', "확률가중 (주관 가정)")]

    def card_row(r0, cards_):
        for k, (lab, val, fm, subt) in enumerate(cards_):
            c1 = B + k * 3; c2 = c1 + 2
            mw(r0, c1, c2, lab, f=font(size=9, bold=True, color="595959"), fill=PALE, al=Alignment(horizontal="center", vertical="center"))
            mw(r0 + 1, c1, c2, val, f=font(size=16, bold=True, color=NAVY), fill=PALE, al=Alignment(horizontal="center", vertical="center"), fmt=fm)
            mw(r0 + 2, c1, c2, subt, f=font(size=8.5, color="7F7F7F"), fill=PALE, al=Alignment(horizontal="center", vertical="center"))
        cs.row_dimensions[r0].height = 18; cs.row_dimensions[r0 + 1].height = 30; cs.row_dimensions[r0 + 2].height = 16
    card_row(KPI_R - 1, cards); card_row(KPI_R + 2, cards2)

    for (k, t_), rr in zip(SECTIONS, TOC):
        c = mw(rr, B, 8, f"   {NUM[k]}.  {t_}", f=Font(name=F, size=10.5, color="0563C1", underline="single"), al=Alignment(vertical="center"))
        c.hyperlink = f"#'{SN}'!B{SEC[k]}"; cs.row_dimensions[rr].height = 17
    cs.freeze_panes = "A3"
    cs.page_setup.orientation = "portrait"; cs.page_setup.fitToWidth = 1; cs.sheet_properties.pageSetUpPr.fitToPage = True; cs.page_setup.fitToHeight = 0


def default_out_dir():
    """보고서 저장 폴더: <프로젝트 루트>/apt_saramara. 프로젝트 루트 = 현재 폴더에서 위로 올라가며 .claude 를 가진 첫 폴더
    (스킬이 .claude/skills/ 안에 있으면 그 .claude 의 부모). 못 찾으면 현재 폴더. 환경변수 APT_SARAMARA_OUT 이 있으면 그것을 우선."""
    import os
    if os.environ.get("APT_SARAMARA_OUT"):
        return Path(os.environ["APT_SARAMARA_OUT"])
    for base in (Path.cwd(), Path(__file__).resolve().parent):
        for d in (base, *base.parents):
            if d.name == ".claude":
                return d.parent / "apt_saramara"
            if (d / ".claude").is_dir():
                return d / "apt_saramara"
    return Path.cwd() / "apt_saramara"


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="아파트 단지 분석 엑셀 보고서 생성")
    ap.add_argument("--apt", required=True, help="config/apartments/<단지>.json")
    ap.add_argument("--market", default=str(Path(__file__).resolve().parent.parent / "config" / "market.json"))
    ap.add_argument("--out", default=None, help="저장 경로. 생략하면 <프로젝트 루트>/apt_saramara/<단지>_분석보고서_<날짜>.xlsx")
    a = ap.parse_args()
    apt = load_json(a.apt); mkt = load_json(a.market)
    if a.out:
        out = Path(a.out)
    else:
        from datetime import date as _d
        out = default_out_dir() / f"{apt['name']}_분석보고서_{_d.today().isoformat()}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    p, M = build(a.apt, a.market, str(out))
    print(f"생성 완료: {Path(p).resolve()}")
